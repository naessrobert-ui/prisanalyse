#!/usr/bin/env python3
"""Les mobile.de i en synlig nettleser. Python 3.10+. Se LES_MEG.md."""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import time
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import parse_qs, urljoin, urlsplit

DEFAULT_SEARCH = "https://suchen.mobile.de/fahrzeuge/search.html?isSearchRequest=true&s=Car&vc=Car&dam=false&ms=13200%3B52&ml=%3A125000&p=%3A27500&vat=1&ref=dsp"
BASE = Path(__file__).resolve().parent

# Tolk bare navngitte bilfelter. Ingen modellbaserte vekt-/batteriestimater.
LABELS = {
    "km": ["Kilometerstand", "Mileage"],
    "forstegangsregistrering": ["Erstzulassung", "First registration"],
    "modellaar": ["Modelljahr", "Model year"],
    "effekt_kw": ["Leistung", "Power"],
    "drivstoff": ["Kraftstoffart", "Fuel type", "Anderer Energieträger"],
    "drivlinje": ["Antriebsart", "Drive type"],
    "girkasse": ["Getriebe", "Transmission"],
    "batteri_kwh": ["Batteriekapazität (in kWh)", "Batteriekapazität", "Battery capacity (in kWh)", "Battery capacity"],
    "rekkevidde_wltp_km": ["Reichweite (WLTP)", "Range (WLTP)", "Elektrische Reichweite (EAER)", "Elektrische Reichweite"],
    "egenvekt_kg": ["Gewicht", "Leergewicht", "Unladen weight", "Weight"],
    "co2_g_km": ["CO₂-Emissionen (komb.)", "CO2-Emissionen (komb.)", "CO₂ emissions (comb.)"],
    "farge": ["Farbe (Hersteller)", "Farbe", "Colour", "Color"],
    "tilstand": ["Fahrzeugzustand", "Vehicle condition"],
    "kategori": ["Kategorie", "Category"],
    "antall_eiere": ["Anzahl der Fahrzeughalter", "Fahrzeughalter", "Number of vehicle owners"],
    "seter": ["Anzahl Sitzplätze", "Number of seats"],
    "dorer": ["Anzahl der Türen", "Number of doors"],
    "eu_kontroll": ["HU", "HU/AU", "Roadworthiness"],
    "interior": ["Innenausstattung", "Interior design"],
    "variant": ["Ausstattungslinie", "Trim line"],
    "serie": ["Baureihe", "Model range"],
    "forhandler_ref": ["Fahrzeugnummer", "Vehicle number"],
    "vin": ["FIN", "Fahrzeug-Identifizierungsnummer", "Vehicle identification number", "VIN"],
}

COLUMNS = {
    "annonse_id": "Annonse-ID", "tittel": "Bil / annonsetittel", "merke": "Merke",
    "modell": "Modell", "variant": "Variant", "pris_eur": "Annonsert bilpris EUR",
    "nettopris_eur": "Oppgitt nettopris EUR", "mva_oppgitt": "Mva-status i annonsen",
    "mva_prosent": "Oppgitt mva %", "km": "Kilometerstand",
    "forstegangsregistrering": "Førstegangsregistrering", "modellaar": "Oppgitt modellår",
    "drivstoff": "Drivstoff", "drivlinje": "Drivlinje", "girkasse": "Girkasse",
    "effekt_kw": "Effekt kW", "effekt_hk": "Effekt hk (PS)",
    "batteri_kwh": "Batteri kWh", "rekkevidde_wltp_km": "Rekkevidde WLTP km",
    "egenvekt_kg": "Oppgitt egenvekt kg", "co2_g_km": "CO2 g/km", "farge": "Farge",
    "tilstand": "Tilstand", "kategori": "Kategori", "antall_eiere": "Antall eiere",
    "seter": "Seter", "dorer": "Dører", "eu_kontroll": "EU-kontroll (HU)",
    "interior": "Interiør", "serie": "Modellserie", "forhandler_ref": "Forhandlerreferanse",
    "vin": "VIN hvis oppgitt", "selger": "Selger", "sted": "Sted",
    "utstyr": "Oppgitt utstyr", "bilde_url": "Hovedbilde URL", "url": "Annonselenke",
    "hentet_utc": "Hentet UTC", "status": "Datastatus", "merknad": "Merknad",
}

# DOM-lesing bruker semantiske felter og synlig tekst, ikke skjulte API-er.
SNAPSHOT_JS = r"""() => {
  const visible = el => !!(el && el.getClientRects().length);
  const text = el => el ? (el.innerText || el.textContent || '').trim() : '';
  const pairs = [];
  for (const dt of document.querySelectorAll('dt')) {
    let dd = dt.nextElementSibling;
    if (dd && dd.tagName === 'DD' && visible(dt)) pairs.push([text(dt), text(dd)]);
  }
  for (const tr of document.querySelectorAll('tr')) {
    const cells = tr.querySelectorAll('th,td');
    if (cells.length === 2 && visible(tr)) pairs.push([text(cells[0]),text(cells[1])]);
  }
  const jsonld = [];
  for (const s of document.querySelectorAll('script[type="application/ld+json"]')) {
    try { jsonld.push(JSON.parse(s.textContent)); } catch (_) {}
  }
  const main = document.querySelector('main') || document.body;
  const headings = [...main.querySelectorAll('h1,h2')].filter(visible).map(text);
  const metas = Object.fromEntries([...document.querySelectorAll('meta[property],meta[name]')]
    .map(m=>[m.getAttribute('property') || m.name, m.content]));
  const prices = [...document.querySelectorAll('[data-testid*="price" i],[data-test*="price" i],[class*="price" i]')]
    .filter(e=>visible(e) && text(e).length < 160).map(text);
  return {url:location.href, title:document.title, headings, pairs, jsonld, metas, prices,
          text:main.innerText, captured_utc:new Date().toISOString()};
}"""

SEARCH_JS = r"""() => {
  const main = document.querySelector('main') || document.body;
  const selectors = '[data-testid="result-listing"],[data-testid="result-list-item"],'
    + '[data-testid="search-result"],[data-testid="search-result-item"],[data-result-item]';
  const cards = [...main.querySelectorAll(selectors)];
  const scope = cards.length ? cards : [main];
  const links = [...new Set(scope.flatMap(root=>[
    ...(root.matches('a[href]') ? [root] : []), ...root.querySelectorAll('a[href]')])
    .filter(a=>a.getClientRects().length).map(a=>a.href)
    .filter(h=>/\/fahrzeuge\/details\.html\?/.test(h)))];
  return {links, scoped:cards.length > 0};
}"""


def now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def label_key(value):
    value = unicodedata.normalize("NFKC", clean(value)).casefold()
    return re.sub(r"[\s:*\d]+$", "", value)


def number(value, machine=False):
    """Tyske grupper/komma; JSON-numerikk med punktum behandles separat."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    match = re.search(r"\d[\d.,\s\u00a0\u202f]*", str(value))
    if not match:
        return None
    token = re.sub(r"\s", "", match.group()).rstrip(".,")
    if "," in token and "." in token:
        if token.rfind(",") > token.rfind("."):
            token = token.replace(".", "").replace(",", ".")
        else:
            token = token.replace(",", "")
    elif "," in token:
        token = token.replace(",", "" if machine and re.fullmatch(r"\d{1,3}(,\d{3})+", token) else ".")
    elif not machine and re.fullmatch(r"\d{1,3}(\.\d{3})+", token):
        token = token.replace(".", "")
    try:
        result = float(token)
        return int(result) if result.is_integer() else result
    except ValueError:
        return None


def canonical_ad(url):
    parts = urlsplit(url)
    if parts.scheme not in ("http", "https") or parts.hostname not in {"suchen.mobile.de", "m.mobile.de", "www.mobile.de"}:
        return None
    ident = parse_qs(parts.query).get("id", [""])[0]
    if parts.path != "/fahrzeuge/details.html" or not ident.isdigit():
        return None
    return "https://suchen.mobile.de/fahrzeuge/details.html?id=" + ident


def valid_search(url):
    parts = urlsplit(url)
    return parts.scheme == "https" and parts.hostname in {"suchen.mobile.de", "m.mobile.de", "www.mobile.de"} and parts.path == "/fahrzeuge/search.html"


def walks(value):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from walks(child)
    elif isinstance(value, list):
        for child in value:
            yield from walks(child)


def named(value):
    return clean(value.get("name", "")) if isinstance(value, dict) else clean(value)


def schema_vehicle(snap, ident):
    candidates = []
    for obj in walks(snap.get("jsonld", [])):
        kinds = obj.get("@type", [])
        kinds = [kinds] if isinstance(kinds, str) else kinds
        if not any(k in {"Car", "Vehicle", "Product"} for k in kinds):
            continue
        reference = str(obj.get("url", obj.get("@id", "")))
        linked_ad = canonical_ad(reference)
        if linked_ad and parse_qs(urlsplit(linked_ad).query)["id"][0] != ident:
            continue  # Ingen anbefalte biler fra JSON-LD.
        score = (10 if ident and ident in reference else 0) + (3 if "Car" in kinds or "Vehicle" in kinds else 0)
        candidates.append((score, obj))
    if not candidates:
        return {}
    matching = [item for item in candidates if item[0] >= 10]
    if matching:
        return max(matching, key=lambda item:item[0])[1]
    # Flere produkter uten identifiserbar annonse-ID kan være anbefalinger.
    return candidates[0][1] if len(candidates) == 1 else {}


def field_values(snap):
    values = {}
    for label, val in snap.get("pairs", []):
        if clean(val):
            values.setdefault(label_key(label), clean(val))
    lines = [clean(line) for line in snap.get("text", "").splitlines() if clean(line)]
    aliases = {label_key(alias): field for field, names in LABELS.items() for alias in names}
    for i, line in enumerate(lines):
        key = label_key(line)
        if key in aliases and i + 1 < len(lines) and label_key(lines[i + 1]) not in aliases:
            values.setdefault(key, lines[i + 1])
        for alias in aliases:
            if key.startswith(alias + ":"):
                values.setdefault(alias, line.split(":", 1)[1].strip())
    return {field: next((values[label_key(a)] for a in names if values.get(label_key(a))), "")
            for field, names in LABELS.items()}


def prices_from_text(snap):
    # Bare første hovedprisblokk før teknikk/finansieringsberegninger og anbefalinger.
    body = snap.get("text", "")
    boundary = re.search(r"(?im)^\s*(Kilometerstand|Mileage|Technische Daten|Technical data|Weitere Fahrzeuge|Ähnliche Fahrzeuge)\s*$", body)
    top = body[:boundary.start()] if boundary else body[:3500]
    gross, net, vat = None, None, None
    candidates = top.splitlines()
    # Priselementer er kun reserve når toppen ikke har en entydig kjøpspris.
    candidates += [p for p in snap.get("prices", []) if clean(p) in clean(top)]
    money = re.compile(r"([\d][\d.\s\u00a0,]*?)\s*(?:€|EUR)", re.I)
    for line in candidates:
        low = line.casefold()
        if re.search(r"mtl|monat|month|rate|anzahlung|darlehen|leasing|finanzier|/jahr|/l\b|/kwh|\bab\s*\d", low):
            continue
        matches = list(money.finditer(line))
        for index, match in enumerate(matches):
            amount = number(match.group(1))
            if amount is None:
                continue
            end = matches[index+1].start() if index+1 < len(matches) else len(line)
            tail = line[match.end():end].casefold()
            prefix = line[max(0, match.start()-18):match.start()].casefold()
            if "netto" in tail or re.search(r"netto\s*:?\s*$", prefix):
                if net is None:
                    net = amount
            elif gross is None:
                gross = amount
        m = re.search(r"(\d{1,2}(?:[.,]\d+)?)\s*%\s*(?:MwSt|VAT)", line, re.I)
        if m:
            vat = number(m.group(1))
    return gross, net, vat, top


def parse_ad(snap, url):
    url = canonical_ad(url)
    if not url:
        raise ValueError("Ugyldig mobile.de-annonselenke")
    ident = parse_qs(urlsplit(url).query)["id"][0]
    row = {key: None for key in COLUMNS}
    sources, warnings = {}, []
    row.update(annonse_id=ident, url=url, hentet_utc=snap.get("captured_utc") or now())
    raw = field_values(snap)
    vehicle = schema_vehicle(snap, ident)
    meta = snap.get("metas", {})
    headings = snap.get("headings", [])
    title = named(vehicle.get("name")) or meta.get("og:title", "") or (headings[0] if headings else "")
    if not title:
        title = snap.get("title", "").split(" - mobile.de")[0]
    row["tittel"] = clean(title)
    row["merke"] = named(vehicle.get("brand")) or None
    row["modell"] = named(vehicle.get("model")) or None
    for field, value in raw.items():
        if value:
            row[field] = value
            sources[field] = value
    for field in ["km", "modellaar", "effekt_kw", "batteri_kwh", "rekkevidde_wltp_km", "egenvekt_kg", "co2_g_km", "antall_eiere", "seter"]:
        row[field] = number(raw[field])
    if raw["effekt_kw"]:
        kw = re.search(r"([\d.,]+)\s*kW\b", raw["effekt_kw"], re.I)
        hp = re.search(r"([\d.,]+)\s*(?:PS|hp)\b", raw["effekt_kw"], re.I)
        row["effekt_kw"] = number(kw.group(1)) if kw else None
        row["effekt_hk"] = number(hp.group(1)) if hp else None
    # JSON-LD-reserver med enhetskontroll.
    odo = vehicle.get("mileageFromOdometer", {})
    if row["km"] is None and isinstance(odo, dict):
        if str(odo.get("unitCode", odo.get("unitText", ""))).casefold() in {"kmt", "km", "kilometer", "kilometre"}:
            row["km"] = number(odo.get("value"), machine=True)
            sources["km"] = "JSON-LD mileageFromOdometer (km)"
    for field, key in [("forstegangsregistrering", "dateVehicleFirstRegistered"), ("modellaar", "vehicleModelDate"), ("farge", "color"), ("vin", "vehicleIdentificationNumber"), ("drivstoff", "fuelType"), ("girkasse", "vehicleTransmission")]:
        if row.get(field) in (None, "") and vehicle.get(key):
            row[field] = number(vehicle[key]) if field == "modellaar" else named(vehicle[key])
            sources[field] = "JSON-LD " + key
    gross, net, vat, top = prices_from_text(snap)
    row.update(pris_eur=gross, nettopris_eur=net, mva_prosent=vat)
    if gross is not None:
        sources["pris_eur"] = "Synlig kjøpspris øverst i annonsen"
    # Tilbudspris i JSON-LD brukes kun med eksplisitt EUR; ingen lowPrice/avdrag.
    offers = vehicle.get("offers", [])
    offers = [offers] if isinstance(offers, dict) else offers
    for offer in offers if isinstance(offers, list) else []:
        if row["pris_eur"] is None and isinstance(offer, dict) and offer.get("priceCurrency") == "EUR" and "price" in offer:
            offer_text = json.dumps(offer, ensure_ascii=False).casefold()
            if any(word in offer_text for word in ("leaseout", "leasing", "monthly", "unitpric", "billingduration", "monat")):
                continue
            amount = number(offer["price"], machine=True)
            if amount is not None and amount >= 1000:
                row["pris_eur"] = amount
                sources["pris_eur"] = "JSON-LD offers.price EUR (prisgrunnlag bør kontrolleres)"
    low = top.casefold()
    if re.search(r"(?:mwst\.?|vat)\s*(?:ist\s*)?(?:nicht|not)\s*(?:ausweisbar|deductible)|differenzbesteuer", low):
        row["mva_oppgitt"] = "Ikke spesifisert/fradragsberettiget ifølge annonsetekst"
    elif net is not None or re.search(r"(?:mwst\.?|vat)\s*(?:ist\s*)?(?:ausweisbar|deductible)", low):
        row["mva_oppgitt"] = "Spesifisert i annonsen – eksportvilkår må bekreftes"
    else:
        row["mva_oppgitt"] = "Ukjent"
    seller = vehicle.get("seller", {})
    if not seller:
        seller = next((o.get("seller") for o in offers if isinstance(o, dict) and o.get("seller")), {}) if isinstance(offers, list) else {}
    if isinstance(seller, dict):
        row["selger"] = named(seller.get("name")) or None
        address = seller.get("address", {})
        if isinstance(address, dict):
            row["sted"] = " ".join(filter(None, [named(address.get("addressCountry")), clean(address.get("postalCode")), clean(address.get("addressLocality"))])) or None
    if not row["sted"]:
        location = re.search(r"(?m)^\s*([A-Z]{2}-\d{4,6}\s+[^\n]+)", top)
        if location:
            row["sted"] = clean(location.group(1))
    if not row["selger"]:
        # På mobile vises firmanavn normalt før stjerner og poststed.
        lines = [clean(x) for x in top.splitlines() if clean(x)]
        for i, line in enumerate(lines):
            if re.match(r"[A-Z]{2}-\d{4,6}\b", line):
                preceding = lines[max(0, i-3):i]
                for candidate in reversed(preceding):
                    if not re.search(r"stern|star|bewertung|^\d[\d.,]*$", candidate, re.I):
                        row["selger"] = candidate
                        break
                break
    image = vehicle.get("image") or meta.get("og:image")
    if isinstance(image, list):
        image = image[0] if image else None
    if isinstance(image, dict):
        image = image.get("url") or image.get("contentUrl")
    row["bilde_url"] = image if isinstance(image, str) and image.startswith("https://") else None
    equipment = re.search(r"(?ims)^\s*(?:Ausstattung|Features|Equipment)\s*$\n(.*?)(?=^\s*(?:Fahrzeugbeschreibung|Vehicle description|Anbieter|Seller|Finanzierung|Preis)\s*$|\Z)", snap.get("text", ""))
    if equipment:
        row["utstyr"] = clean(equipment.group(1))[:12000]
    missing = [COLUMNS[k] for k in ("tittel", "pris_eur", "km", "forstegangsregistrering") if row[k] in (None, "")]
    if missing:
        warnings.append("Mangler: " + ", ".join(missing))
    if row["egenvekt_kg"] is None:
        warnings.append("Egenvekt ikke oppgitt/ikke funnet; ingen estimering")
    row["status"] = "ok" if not missing else "ufullstendig"
    row["merknad"] = "; ".join(warnings)
    row["kilder"] = sources
    row["raa_tekniske_felt"] = raw
    return row


def safe_cell(value):
    # Forhandlertekst behandles som tekst også når den begynner med =, +, - eller @.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def atomic_json(path, value):
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def save_results(folder, state, excel=False):
    folder.mkdir(parents=True, exist_ok=True)
    state["sist_lagret_utc"] = now()
    atomic_json(folder / "biler.json", state)
    temporary = folder / "biler.csv.tmp"
    with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(COLUMNS.values())
        for row in state["biler"]:
            values = [safe_cell(row.get(key)) for key in COLUMNS]
            writer.writerow([str(v).replace(".", ",") if isinstance(v, float) else v for v in values])
    temporary.replace(folder / "biler.csv")
    if not excel:
        return
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    book = Workbook()
    sheet = book.active
    sheet.title = "Biler"
    sheet.append(list(COLUMNS.values()))
    for row in state["biler"]:
        sheet.append([safe_cell(row.get(k)) for k in COLUMNS])
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173D4E")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 34
    for index, key in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 35 if key in {"tittel", "merknad", "url", "mva_oppgitt"} else 21
        for cells in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
            cell = cells[0]
            if key in {"pris_eur", "nettopris_eur", "km", "egenvekt_kg"}:
                cell.number_format = '#,##0.00' if isinstance(cell.value, float) else '#,##0'
            if key == "url" and cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"
    info = book.create_sheet("Om uttrekket")
    for key in ("sok_url", "startet_utc", "sist_lagret_utc", "sok_status", "sok_sider", "sok_merknader", "antall_lenker", "kjorestatus", "feil"):
        value = state.get(key)
        info.append([key, safe_cell(json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value)])
    info.append(["Prisgrunnlag", "Bilpris i EUR slik annonsen oppgir den. Norsk frakt, avgifter og registrering er ikke beregnet."])
    info.append(["Nettopris", "Bare eksplisitt oppgitt nettopris; mva/eksportmulighet må bekreftes med selger."])
    info.append(["Tomt felt", "Ikke oppgitt eller ikke funnet. Ikke lik null. Ingen vekt-/batteriestimater."])
    info.append(["Dato", "Førstegangsregistrering og modellår er forskjellige felter."])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 110
    for row in info:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    temp_xlsx = folder / "biler.tmp.xlsx"
    book.save(temp_xlsx)
    temp_xlsx.replace(folder / "biler.xlsx")


class AccessStop(RuntimeError):
    pass


def navigate(page, url, delay):
    time.sleep(delay)
    response = page.goto(url, wait_until="domcontentloaded", timeout=60000)
    if response and response.status in {401, 403, 429}:
        raise AccessStop(f"mobile.de svarte HTTP {response.status}. Uttrekket stoppes; data som er hentet beholdes.")
    if response and response.status >= 400:
        raise RuntimeError(f"HTTP {response.status}: {url}")
    page.wait_for_timeout(2200)
    body = page.locator("body").inner_text(timeout=15000)
    if re.search(r"access denied|zugriff verweigert|too many requests|unusual traffic|ungewöhnlich(?:e|en) zugriffe", body, re.I):
        raise AccessStop("mobile.de viser en tilgangsbegrensning. Uttrekket stoppes.")
    if re.search(r"verify (?:that )?you are human|bestätigen sie.*mensch|ich bin kein roboter|security check|sicherheitsüberprüfung", body, re.I):
        input("Nettleseren viser en kontroll. Fullfør den manuelt hvis mulig, og trykk Enter her. Ctrl+C avbryter. ")
        body = page.locator("body").inner_text()
        if re.search(r"verify (?:that )?you are human|ich bin kein roboter|security check|sicherheitsüberprüfung", body, re.I):
            raise AccessStop("Kontrollen er fortsatt synlig. Uttrekket stoppes.")


def next_search_url(page):
    # Bare faktiske neste-lenker; gjett aldri sidetall og avslutt aldri ved feil.
    selectors = ['a[rel="next"]', 'a[aria-label*="Nächste" i]', 'a[aria-label*="Next" i]',
                 '[data-testid*="pagination" i] a[aria-label*="weiter" i]',
                 'a[data-testid*="pagination-next" i]']
    for selector in selectors:
        for element in page.locator(selector).all():
            if not element.is_visible() or element.get_attribute("aria-disabled") == "true":
                continue
            href = element.get_attribute("href")
            if href:
                candidate = urljoin(page.url, href)
                if valid_search(candidate):
                    return candidate
    for element in page.get_by_role("link", name=re.compile(r"^(?:Nächste(?: Seite)?|Weiter|Next(?: page)?)\s*[›»→]?$", re.I)).all():
        if element.is_visible() and element.get_attribute("aria-disabled") != "true":
            href = element.get_attribute("href")
            if href and valid_search(urljoin(page.url, href)):
                return urljoin(page.url, href)
    return None


def next_search_button(page):
    for selector in ('button[aria-label*="Nächste" i]', 'button[aria-label*="Next" i]',
                     'button[data-testid*="pagination-next" i]', '[data-testid*="pagination-next" i] button'):
        for button in page.locator(selector).all():
            if button.is_visible() and button.is_enabled() and button.get_attribute("aria-disabled") != "true":
                return button
    for button in page.get_by_role("button", name=re.compile(r"^(?:Nächste(?: Seite)?|Weiter|Next(?: page)?)\s*[›»→]?$", re.I)).all():
        if button.is_visible() and button.is_enabled() and button.get_attribute("aria-disabled") != "true":
            return button
    return None


def collect_links(page, args, state, folder):
    seen_signatures = set()
    urls = list(state.get("annonselenker", []))
    seen_ads = set(urls)
    for count in range(1, args.max_pages + 1):
        for _ in range(3):
            page.mouse.wheel(0, 1600)
            page.wait_for_timeout(600)
        page.evaluate("window.scrollTo(0,0)")
        snapshot = page.evaluate(SEARCH_JS)
        links = list(dict.fromkeys(filter(None, (canonical_ad(u) for u in snapshot["links"]))))
        if not links:
            state["sok_status"] = "ingen annonselenker funnet: kontroller nettsiden"
            break
        signature = tuple(sorted(links))
        if signature in seen_signatures:
            state["sok_status"] = "stoppet: samme annonser på neste side"
            break
        seen_signatures.add(signature)
        if not snapshot["scoped"]:
            note = "Lenker lest fra sidens hovedinnhold; anbefalte annonser kan være med. Kontroller treffene mot søket."
            if note not in state["sok_merknader"]:
                state["sok_merknader"].append(note)
        for url in links:
            if url not in seen_ads:
                seen_ads.add(url)
                urls.append(url)
        state.update(sok_sider=count, annonselenker=urls[:args.max_cars], antall_lenker=min(len(urls), args.max_cars))
        atomic_json(folder / "biler.json", state)
        print(f"Søkeside {count}: {len(links)} lenker; {len(urls)} unike hittil.")
        if len(urls) >= args.max_cars:
            state["sok_status"] = f"grense nådd: {args.max_cars} biler; flere kan finnes"
            break
        next_url = next_search_url(page)
        next_button = next_search_button(page) if not next_url else None
        if not next_url and next_button is None:
            state["sok_status"] = "ingen neste-lenke/-knapp funnet; kontroller at alle sider er med"
            break
        if count == args.max_pages:
            state["sok_status"] = f"grense nådd: {args.max_pages} søkesider; flere finnes"
            break
        if next_url:
            navigate(page, next_url, args.delay)
        else:
            time.sleep(args.delay)
            next_button.click(timeout=10000)
            page.wait_for_timeout(2500)
            if re.search(r"access denied|zugriff verweigert|too many requests|unusual traffic", page.locator("body").inner_text(), re.I):
                raise AccessStop("mobile.de viser en tilgangsbegrensning på neste søkeside.")
    return urls[:args.max_cars]


def open_context(playwright, args):
    options = {"user_data_dir":str(BASE / "nettleserprofil"), "headless":False,
               "locale":"de-DE", "viewport":{"width":1360, "height":960}}
    channels = [args.browser] if args.browser != "auto" else ["chrome", "msedge", "chromium"]
    errors = []
    for channel in channels:
        try:
            return playwright.chromium.launch_persistent_context(**options, **({"channel":channel} if channel != "chromium" else {}))
        except Exception as exc:
            errors.append(f"{channel}: {str(exc).splitlines()[0]}")
    raise RuntimeError("Kunne ikke åpne nettleser. Lukk andre kjøringer av skriptet. Installer Chrome/Edge, eller kjør: python -m playwright install chromium. " + " | ".join(errors))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--url", help="mobile.de-søk; standard er lenken Robert oppga")
    parser.add_argument("--links", type=Path, help="Tekstfil med én mobile.de-annonselenke per linje")
    parser.add_argument("--output", type=Path, help="Ny resultatmappe; standard: resultater/dato-klokkeslett")
    parser.add_argument("--resume", type=Path, help="Fortsett en eksisterende resultatmappe og prøv ufullstendige/feilede annonser igjen")
    parser.add_argument("--max-pages", type=int, default=50)
    parser.add_argument("--max-cars", type=int, default=1000)
    parser.add_argument("--delay", type=float, default=3, help="Sekunder mellom sideåpninger, minst 2")
    parser.add_argument("--browser", choices=["auto", "chrome", "msedge", "chromium"], default="auto")
    args = parser.parse_args(argv)
    if args.max_pages < 1 or args.max_cars < 1 or args.delay < 2:
        parser.error("max-pages og max-cars må være positive, delay minst 2 sekunder")
    if args.resume and (args.output or args.links):
        parser.error("--resume kan ikke kombineres med --output eller --links")
    search = args.url or DEFAULT_SEARCH
    if not valid_search(search):
        parser.error("--url må være en HTTPS-søkelenke på mobile.de")
    if args.resume:
        folder = args.resume.resolve()
        state = json.loads((folder / "biler.json").read_text(encoding="utf-8"))
        if args.url and state["sok_url"] != args.url:
            parser.error("Søkelenken stemmer ikke med uttrekket som skal fortsettes")
        search = state["sok_url"]
    else:
        folder = (args.output or BASE / "resultater" / datetime.now().strftime("%Y-%m-%d_%H-%M-%S")).resolve()
        if (folder / "biler.json").exists():
            parser.error("Mappen har allerede data. Bruk --resume eller en ny --output-mappe")
        state = {"formatversjon":1, "sok_url":search, "startet_utc":now(), "biler":[],
                 "annonselenker":[], "sok_sider":0, "antall_lenker":0, "sok_status":"ikke startet",
                 "sok_merknader":[], "kjorestatus":"pågår", "feil":[]}
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "kildedata").mkdir(exist_ok=True)
    if args.links:
        urls = []
        for line in args.links.read_text(encoding="utf-8-sig").splitlines():
            if not line.strip() or line.lstrip().startswith("#"):
                continue
            link = canonical_ad(line.strip())
            if not link:
                parser.error("Ugyldig annonselenke i --links: " + line)
            if link not in urls:
                urls.append(link)
        if not urls:
            parser.error("Lenkefilen inneholder ingen annonselenker")
        state.update(annonselenker=urls[:args.max_cars], antall_lenker=min(len(urls), args.max_cars),
                     sok_status="manuell lenkeliste" + ("; bilgrense nådd" if len(urls)>args.max_cars else ""))
    try:
        from playwright.sync_api import sync_playwright
        import openpyxl  # Avklar avhengigheter før nettleseren starter.
    except ImportError:
        print("Installer først: python -m pip install -r requirements.txt", file=sys.stderr)
        return 1
    context = None
    exit_code = 0
    try:
        with sync_playwright() as pw:
            try:
                context = open_context(pw, args)
                page = context.pages[0] if context.pages else context.new_page()
                start_url = state["annonselenker"][0] if args.links or (args.resume and state["annonselenker"]) else search
                navigate(page, start_url, 0)
                print("\nSøket/annonsen er åpnet i nettleseren. Velg informasjonskapsler selv.")
                input("Når bilene/annonsen er synlig: trykk Enter HER for å starte. Ctrl+C avbryter. ")
                if not args.links and not (args.resume and state["annonselenker"]):
                    collect_links(page, args, state, folder)
                done = {r["annonse_id"] for r in state["biler"] if r.get("status") == "ok"}
                urls = state["annonselenker"]
                if not urls:
                    raise RuntimeError("Ingen annonselenker ble funnet. Kontroller søket eller bruk --links med innlimte annonselenker.")
                for i, url in enumerate(urls, 1):
                    ident = parse_qs(urlsplit(url).query)["id"][0]
                    if ident in done:
                        continue
                    print(f"[{i}/{len(urls)}] Leser annonse {ident} ...", flush=True)
                    try:
                        navigate(page, url, args.delay)
                        # Åpne bare eksplisitte teknikk-/utstyrsknapper, aldri kontakt/kjøp.
                        for label in ("Weitere technische Daten", "Alle technischen Daten", "Alle Ausstattungen anzeigen"):
                            button = page.get_by_role("button", name=label, exact=True)
                            if button.count() == 1 and button.is_visible():
                                button.click(timeout=3000)
                                page.wait_for_timeout(300)
                        snapshot = page.evaluate(SNAPSHOT_JS)
                        atomic_json(folder / "kildedata" / f"{ident}.json", snapshot)
                        final_url = canonical_ad(snapshot["url"])
                        if final_url != url:
                            raise RuntimeError("Annonsen ble omdirigert til en annen side; ingen bildata godtatt")
                        row = parse_ad(snapshot, url)
                    except AccessStop:
                        raise
                    except Exception as exc:
                        row = {"annonse_id":ident, "url":url, "hentet_utc":now(), "status":"feil", "merknad":str(exc)}
                    state["biler"] = [r for r in state["biler"] if r["annonse_id"] != ident] + [row]
                    save_results(folder, state)
                    print("  " + (row.get("tittel") or ident) + " | " + row["status"])
                state["kjorestatus"] = "lenkeliste ferdig behandlet"
            finally:
                if context:
                    context.close()
    except KeyboardInterrupt:
        state["kjorestatus"] = "avbrutt av bruker"
        exit_code = 130
    except Exception as exc:
        state["kjorestatus"] = "stoppet med feil"
        state["feil"].append({"tid":now(), "melding":str(exc)})
        print(str(exc), file=sys.stderr)
        exit_code = 1
    finally:
        try:
            save_results(folder, state, excel=True)
        except Exception as exc:
            print("Kunne ikke skrive alle eksportfiler (lukk Excel hvis filen er åpen): " + str(exc), file=sys.stderr)
            exit_code = 1
    good = sum(r.get("status") == "ok" for r in state["biler"])
    print(f"\n{good} annonser med kjernefelter, {len(state['biler'])-good} ufullstendige/feil.")
    print("Søkeomfang: " + state["sok_status"])
    for note in state["sok_merknader"]:
        print(note)
    print("Resultatmappe: " + str(folder))
    print('Fortsett ved behov: python mobile_hent.py --resume "' + str(folder) + '"')
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())
