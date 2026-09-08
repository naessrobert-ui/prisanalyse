import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook
import import_radar_routes as routes
from import_radar_search import SourceError
from mobile_search_export import DEFAULT_SEARCH, collect, snapshot_html, validate_url
from scripts.mobile_hent import parse_ad

AD = "https://suchen.mobile.de/fahrzeuge/details.html?id=123"
HTML = '''<main><h2 data-testid="vip-ad-title">Testbil Electric</h2>
<div data-testid="main-price-area"><span>27.500</span><span>€</span><span>23.109 € Netto, 19% MwSt.</span></div>
<dl><dt>Kilometerstand</dt><dd>80.000 km</dd><dt>Erstzulassung</dt><dd>03/2022</dd>
<dt>Batteriekapazität</dt><dd>64,8 kWh</dd></dl></main>'''
SEARCH = f'<main><a data-testid="result-listing" href="{AD}">Testbil</a></main>'


class MobileTests(unittest.TestCase):
    def test_validate_rejects_external_credentials_ports_and_nonsearch(self):
        self.assertEqual(validate_url(DEFAULT_SEARCH), DEFAULT_SEARCH)
        for bad in [DEFAULT_SEARCH.replace("https:","http:"), DEFAULT_SEARCH.replace("suchen.mobile.de","evil.test"),
                    DEFAULT_SEARCH.replace("suchen.mobile.de","suchen.mobile.de:444"),
                    DEFAULT_SEARCH.replace("suchen.mobile.de","user@suchen.mobile.de"), AD]:
            with self.subTest(bad=bad), self.assertRaises(ValueError):
                validate_url(bad)

    def test_html_adapter_preserves_gross_net_and_unknown_weight(self):
        row = parse_ad(snapshot_html(HTML, AD), AD)
        self.assertEqual((row["pris_eur"],row["nettopris_eur"],row["km"]),(27500,23109,80000))
        self.assertEqual(row["batteri_kwh"],64.8)
        self.assertIsNone(row["egenvekt_kg"])
        self.assertIsNone(row["modellaar"])
        self.assertEqual(row["status"],"ok")

    def test_pagination_dedup_and_checkpoint(self):
        second = DEFAULT_SEARCH + "&pageNumber=2"
        ad2 = AD.replace("123", "124")
        pages = {DEFAULT_SEARCH: SEARCH + f'<a rel="next" href="{second}">Next</a>',
                 second: SEARCH.replace('</main>', f'<a href="{ad2}" data-testid="result-listing">Bil 2</a></main>'), AD:HTML,ad2:HTML}
        calls,counts=[],[]
        def fetch(url,**kw):
            calls.append(url)
            return pages[url]
        report=collect(DEFAULT_SEARCH,20,fetch=fetch,pause=0,checkpoint=lambda r:counts.append(len(r["biler"])))
        self.assertEqual(len(report["biler"]),2)
        self.assertEqual(calls.count(AD),1)
        self.assertEqual(report["sok_sider"],2)
        self.assertIn(1,counts)
        self.assertEqual(report["kjorestatus"],"ferdig")

    def test_block_stops_and_keeps_partial_data(self):
        ad2=AD.replace("123","124")
        search=SEARCH.replace('</main>',f'<a href="{ad2}">Second</a><a href="{ad2}9">Third</a></main>')
        calls=[]
        def fetch(url,**kw):
            calls.append(url)
            if url==DEFAULT_SEARCH:return search.replace('data-testid="result-listing"','')
            if url==AD:return HTML
            raise SourceError("HTTP 403; stoppet")
        r=collect(DEFAULT_SEARCH,20,fetch=fetch,pause=0)
        self.assertEqual(len(r["biler"]),1)
        self.assertEqual(r["kjorestatus"],"delvis")
        self.assertEqual(len(calls),3)
        self.assertIn("403",r["feil"][0])

    def test_session_csrf_job_and_downloads(self):
        root=Path(__file__).resolve().parents[1]
        app=Flask(__name__,template_folder=str(root/"templates"))
        app.secret_key="test-only"
        app.register_blueprint(routes.import_radar_bp)
        with tempfile.TemporaryDirectory() as temp:
            app.config["IMPORT_RADAR_DB_PATH"]=str(Path(temp)/"jobs.sqlite3")
            client=app.test_client()
            base="/bil/import-radar"
            self.assertEqual(client.get(base+"/mobile").status_code,200)
            with client.session_transaction() as s:csrf=s["import_radar_csrf"]
            self.assertEqual(client.post(base+"/api/mobile",json={"url":DEFAULT_SEARCH}).status_code,403)
            def inline(fn,*args):fn(*args)
            original=collect
            def replay(url,limit,checkpoint):
                return original(url,limit,checkpoint=checkpoint,fetch=lambda u,**k:SEARCH if u==DEFAULT_SEARCH else HTML,pause=0)
            with patch.object(routes._POOL,"submit",side_effect=inline),patch("mobile_search_export.collect",side_effect=replay):
                response=client.post(base+"/api/mobile",json={"url":DEFAULT_SEARCH,"limit":1},headers={"X-CSRF-Token":csrf})
            self.assertEqual(response.status_code,202)
            ident=response.json["id"]
            job=client.get(base+"/api/search/"+ident).json
            self.assertEqual(job["status"],"done")
            self.assertEqual(len(job["report"]["biler"]),1)
            for kind in ("json","csv","xlsx"):
                result=client.get(f"{base}/api/mobile/{ident}/download/{kind}")
                self.assertEqual(result.status_code,200)
                self.assertEqual(result.headers["Cache-Control"],"no-store")
                if kind=="xlsx":
                    book=load_workbook(io.BytesIO(result.data))
                    self.assertEqual(book["Biler"].max_row,2)
                    self.assertEqual(book["Biler"]["F2"].value,27500)
            outsider=app.test_client()
            self.assertEqual(outsider.get(f"{base}/api/mobile/{ident}/download/json").status_code,404)
            self.assertEqual(outsider.get(f"{base}/api/search/{ident}").status_code,404)


if __name__=="__main__":
    unittest.main()
